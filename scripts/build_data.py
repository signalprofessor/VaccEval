"""Build a browser-safe VaccEval aggregate from an event-level workbook."""
from __future__ import annotations
import argparse, json
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path
from openpyxl import load_workbook

GROUPS = {
    "COVID-19": {"U071", "U072"},
    "Influenza": {"J09", "J10", "J11"},
    "RSV": {"J121", "J210", "B974"},
    "Pneumonia": {"J13", "J18", "J189"},
    "Bronchitis": {"J20", "J205", "J21", "J210"},
}
REGION_NAMES = {"ro": "Östergötland", "rk": "Kalmar", "rj": "Jönköping"}

def normalized_code(value: object) -> str:
    return str(value or "").strip().upper().replace(".", "")

def as_date(value: object) -> date | None:
    if isinstance(value, datetime): return value.date()
    if isinstance(value, date): return value
    return None

def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    workbook = load_workbook(args.source, read_only=True, data_only=True)
    rows = workbook.active.iter_rows(values_only=True)
    headers = [str(value) for value in next(rows)]
    index = {name: position for position, name in enumerate(headers)}
    required = {"KontaktID", "ICD10KodNr", "KategoriNr", "Startdatum"}
    missing = required - index.keys()
    if missing: raise ValueError(f"Missing columns: {sorted(missing)}")

    encounters: dict[tuple[date, str, str], set[str]] = defaultdict(set)
    source_rows = 0
    first_date: date | None = None
    last_date: date | None = None
    for row in rows:
        source_rows += 1
        encounter = str(row[index["KontaktID"]] or "").strip()
        event_date = as_date(row[index["Startdatum"]])
        if not encounter or event_date is None: continue
        region_code = encounter.split("_", 1)[0].lower()
        if region_code not in REGION_NAMES: continue
        code = normalized_code(row[index["ICD10KodNr"]])
        category = normalized_code(row[index["KategoriNr"]])
        for group, tokens in GROUPS.items():
            if code in tokens or category in tokens:
                encounters[(event_date, region_code, group)].add(encounter)
        first_date = event_date if first_date is None else min(first_date, event_date)
        last_date = event_date if last_date is None else max(last_date, event_date)
    if first_date is None or last_date is None: raise ValueError("No usable encounters found")

    series = []
    current = first_date
    while current <= last_date:
        for region_code, region_name in REGION_NAMES.items():
            for group in GROUPS:
                series.append({"date": current.isoformat(), "region": region_name,
                               "pathogen": group,
                               "count": len(encounters[(current, region_code, group)])})
        current += timedelta(days=1)
    payload = {
        "meta": {"sourceSnapshot": args.source.stem,
                 "generatedAt": datetime.now().astimezone().isoformat(timespec="seconds"),
                 "dateRange": [first_date.isoformat(), last_date.isoformat()],
                 "sourceRows": source_rows,
                 "measure": "Unique healthcare encounters per day",
                 "privacy": "Only daily aggregate counts; no row-level fields or identifiers.",
                 "countingRule": "One encounter per pathogen group, date, and region."},
        "regions": list(REGION_NAMES.values()), "pathogens": list(GROUPS.keys()), "series": series}
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")))
    print(json.dumps({**payload["meta"], "records": len(series)}, ensure_ascii=False, indent=2))

if __name__ == "__main__": main()
